import unittest
import tempfile
from pathlib import Path
import openpyxl
import pandas as pd

from src.database.models import CommentModel, PostModel
from src.database.db_manager import DatabaseManager
from src.exporter.exporter import DataExporter

class TestCommentGenerations(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.db_path = Path(self.temp_dir.name) / "test_generations.db"
        self.export_dir = Path(self.temp_dir.name) / "exports"
        self.db = DatabaseManager(db_path=self.db_path)
        self.exporter = DataExporter(export_dir=self.export_dir)

        # Seed post
        post = PostModel(
            id="post_gen_1",
            url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="post_author",
            content="Mọi người cho mình xin ý kiến về vấn đề này với ạ?",
            categories=["Đời sống"]
        )
        self.db.upsert_post(post)

    def tearDown(self):
        try:
            self.temp_dir.cleanup()
        except Exception:
            pass

    def test_comment_model_default_generations(self):
        """Verify that CommentModel has f0, f1, f2, f3 and reply_level defaulted properly."""
        c = CommentModel(
            id="c_test_1",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            content="Bình luận thử nghiệm"
        )
        self.assertEqual(c.f0, "")
        self.assertEqual(c.f1, "")
        self.assertEqual(c.f2, "")
        self.assertEqual(c.f3, "")
        self.assertEqual(c.reply_level, 0)
        self.assertEqual(c.is_reply, False)

    def test_upsert_and_retrieve_generations(self):
        """Verify that f0, f1, f2, f3 and reply_level are stored and queried correctly."""
        # F0 Root comment
        c0 = CommentModel(
            id="c_root_1",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="user_a",
            content="Ý kiến của tôi là nên bình tĩnh",
            is_reply=False,
            comment_type_vi="Bình luận gốc",
            f0="",
            f1="",
            f2="",
            f3="",
            reply_level=0
        )
        self.db.upsert_comment(c0)

        # F1 Reply comment
        c1 = CommentModel(
            id="c_reply_f1",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="post_author",
            content="Cảm ơn bạn user_a nhé",
            is_reply=True,
            parent_comment_id="c_root_1",
            comment_type_vi="Bình luận con (F1)",
            f0="Ý kiến của tôi là nên bình tĩnh",
            f1="Cảm ơn bạn user_a nhé",
            f2="",
            f3="",
            reply_level=1
        )
        self.db.upsert_comment(c1)

        # F2 Reply comment
        c2 = CommentModel(
            id="c_reply_f2",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="user_b",
            content="Bình tĩnh sao được bà ơi",
            is_reply=True,
            parent_comment_id="c_reply_f1",
            comment_type_vi="Bình luận con (F2)",
            f0="Ý kiến của tôi là nên bình tĩnh",
            f1="Cảm ơn bạn user_a nhé",
            f2="Bình tĩnh sao được bà ơi",
            f3="",
            reply_level=2
        )
        self.db.upsert_comment(c2)

        # F3 Reply comment
        c3 = CommentModel(
            id="c_reply_f3",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="post_author",
            content="Thôi nhịn đi bạn ơi, cãi nhau làm gì",
            is_reply=True,
            parent_comment_id="c_reply_f2",
            comment_type_vi="Bình luận con (F3)",
            f0="Ý kiến của tôi là nên bình tĩnh",
            f1="Cảm ơn bạn user_a nhé",
            f2="Bình tĩnh sao được bà ơi",
            f3="Thôi nhịn đi bạn ơi, cãi nhau làm gì",
            reply_level=3
        )
        self.db.upsert_comment(c3)

        # Verify query in DB
        df = self.db.get_curated_export_df(filter_status="all")
        self.assertEqual(len(df), 4)

        # Check F0 row: f0 is empty because cmt is root ("nếu cmt cần xử lý là reply")
        row_f0 = df[df["Mã ID"] == "c_root_1"].iloc[0]
        self.assertEqual(row_f0["f0"], "")
        self.assertEqual(row_f0["f1"], "")

        # Check F1 row
        row_f1 = df[df["Mã ID"] == "c_reply_f1"].iloc[0]
        self.assertEqual(row_f1["f0"], "Ý kiến của tôi là nên bình tĩnh")
        self.assertEqual(row_f1["f1"], "Cảm ơn bạn user_a nhé")
        self.assertEqual(row_f1["f2"], "")

        # Check F2 row
        row_f2 = df[df["Mã ID"] == "c_reply_f2"].iloc[0]
        self.assertEqual(row_f2["f0"], "Ý kiến của tôi là nên bình tĩnh")
        self.assertEqual(row_f2["f1"], "Cảm ơn bạn user_a nhé")
        self.assertEqual(row_f2["f2"], "Bình tĩnh sao được bà ơi")
        self.assertEqual(row_f2["f3"], "")

        # Check F3 row
        row_f3 = df[df["Mã ID"] == "c_reply_f3"].iloc[0]
        self.assertEqual(row_f3["f0"], "Ý kiến của tôi là nên bình tĩnh")
        self.assertEqual(row_f3["f1"], "Cảm ơn bạn user_a nhé")
        self.assertEqual(row_f3["f2"], "Bình tĩnh sao được bà ơi")
        self.assertEqual(row_f3["f3"], "Thôi nhịn đi bạn ơi, cãi nhau làm gì")

    def test_backfill_comment_generations(self):
        """Test auto-backfilling f0, f1, f2, f3 for a sequential dialogue flow."""
        # Insert sequential comments without pre-computed f0-f3
        c1 = CommentModel(
            id="seq_1",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="commenter_1",
            content="Con này mất dạy thật sự luôn",
            is_reply=False
        )
        self.db.upsert_comment(c1)

        c2 = CommentModel(
            id="seq_2",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="post_author",
            content="Đúng z bà ơi, em tôi nó láo lắm",
            is_reply=False
        )
        self.db.upsert_comment(c2)

        c3 = CommentModel(
            id="seq_3",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="commenter_1",
            content="Bà cho nó một bài học đi",
            is_reply=False
        )
        self.db.upsert_comment(c3)

        # Run backfill
        count = self.db.backfill_comment_generations()
        self.assertEqual(count, 3)

        df = self.db.get_curated_export_df(filter_status="all")
        r1 = df[df["Mã ID"] == "seq_1"].iloc[0]
        r2 = df[df["Mã ID"] == "seq_2"].iloc[0]
        r3 = df[df["Mã ID"] == "seq_3"].iloc[0]

        # seq_1 is Root
        self.assertEqual(r1["f0"], "")
        # seq_2 is F1 reply to seq_1
        self.assertEqual(r2["f0"], "Con này mất dạy thật sự luôn")
        self.assertEqual(r2["f1"], "Đúng z bà ơi, em tôi nó láo lắm")
        # seq_3 is F2 reply to seq_2
        self.assertEqual(r3["f0"], "Con này mất dạy thật sự luôn")
        self.assertEqual(r3["f1"], "Đúng z bà ơi, em tôi nó láo lắm")
        self.assertEqual(r3["f2"], "Bà cho nó một bài học đi")

    def test_curated_excel_export_contains_f0_f3(self):
        """Verify Excel export includes f0, f1, f2, f3 and applies styling."""
        c = CommentModel(
            id="cmt_excel_1",
            post_id="post_gen_1",
            post_url="https://www.threads.net/@post_author/post/XYZ999",
            author_username="post_author",
            content="Nội dung cmt phản hồi",
            is_reply=True,
            f0="Bình luận gốc mẫu",
            f1="Nội dung cmt phản hồi",
            reply_level=1
        )
        self.db.upsert_comment(c)

        from src.database.db_manager import db_manager
        orig_db = db_manager.db_path
        db_manager.db_path = self.db_path
        try:
            excel_path = self.exporter.export_curated_excel(filter_status="all")
            self.assertTrue(excel_path.exists())

            wb = openpyxl.load_workbook(excel_path)
            ws = wb.active
            headers = [cell.value for cell in ws[1] if cell.value is not None]

            self.assertIn("f0", headers)
            self.assertIn("f1", headers)
            self.assertIn("f2", headers)
            self.assertIn("f3", headers)

            # Check column widths configured
            f0_col_idx = headers.index("f0") + 1
            col_letter = openpyxl.utils.get_column_letter(f0_col_idx)
            self.assertGreaterEqual(ws.column_dimensions[col_letter].width, 35)

            wb.close()
        finally:
            db_manager.db_path = orig_db

if __name__ == "__main__":
    unittest.main()
