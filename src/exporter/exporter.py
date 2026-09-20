import json
from pathlib import Path
from datetime import datetime
from typing import Optional, List
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import EXPORTS_DIR
from src.database.db_manager import db_manager
from src.utils.logger import logger

class DataExporter:
    """Exports collected Threads toxic and general data to Excel, CSV, and JSON."""

    def __init__(self, export_dir: Optional[Path] = None):
        self.export_dir = export_dir or EXPORTS_DIR
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def export_detailed_csv(
        self,
        table_type: str = "comments",  # "comments" or "posts"
        toxic_only: bool = False,       # Default False so user can aggregate all data
        filename: Optional[str] = None,
        limit: Optional[int] = None
    ) -> Path:
        """
        Export a comprehensive, human-readable CSV with Vietnamese column names,
        including author profiles, emojis, scores, images, and reply links.
        Uses UTF-8-BOM (utf-8-sig) for native Excel compatibility without garbled text.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "toxic_only" if toxic_only else "all_detailed"
        out_name = filename or f"threads_{table_type}_{suffix}_{timestamp}.csv"
        out_path = self.export_dir / out_name

        df = db_manager.get_detailed_export_df(table_type=table_type, toxic_only=toxic_only, limit=limit)
        df.to_csv(out_path, index=False, encoding="utf-8-sig")
        logger.info(f"Exported detailed CSV file successfully: {out_path} ({len(df)} rows)")
        return out_path

    def export_comments_csv(
        self,
        filter_status: str = "all",
        filter_comment_type: str = "all",
        include_links: bool = False,
        filename: Optional[str] = None,
        limit: Optional[int] = None,
        columns: Optional[List[str]] = None,
        viet_eng_only: bool = True
    ) -> Path:
        """
        Export comments focused CSV:
        - filter_status: 'bad' (Xấu luôn), 'ambiguous' (Chưa rõ), 'all' (Tất cả để tự tổng hợp)
        - filter_comment_type: 'root' (Chỉ bình luận gốc), 'reply' (Chỉ bình luận con), 'all' (Toàn bộ)
        - include_links: False (Chỉ bình luận & đánh giá), True (Kèm theo link bài viết/bình luận/profile/parent_id)
        - limit: None (Không giới hạn)
        - columns: Danh sách cột cần xuất (Nếu người dùng ẩn bớt cột trên giao diện)
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        status_slug_map = {
            "bad": "xau_luon",
            "ambiguous": "chua_ro_nghi_ngo",
            "clean": "trong_sach",
            "all": "tong_hop"
        }
        type_slug_map = {
            "root": "_goc",
            "reply": "_con_phan_hoi",
            "all": ""
        }
        slug = status_slug_map.get(filter_status, filter_status)
        type_slug = type_slug_map.get(filter_comment_type, "")
        link_suffix = "_kem_link" if include_links else ""
        out_name = filename or f"threads_binh_luan_{slug}{type_slug}{link_suffix}_{timestamp}.csv"
        out_path = self.export_dir / out_name

        df = db_manager.get_comments_export_df(
            filter_status=filter_status,
            filter_comment_type=filter_comment_type,
            include_links=include_links,
            limit=limit,
            viet_eng_only=viet_eng_only
        )
        if columns:
            valid_cols = [c for c in columns if c in df.columns]
            if valid_cols:
                df = df[valid_cols]

        df.to_csv(out_path, index=False, encoding="utf-8-sig")
        logger.info(f"Exported comments CSV successfully: {out_path} ({len(df)} rows)")
        return out_path

    def export_to_csv(
        self,
        table_type: str = "comments",
        toxic_only: bool = True,
        filename: Optional[str] = None,
        detailed: bool = True,
        limit: Optional[int] = None
    ) -> Path:
        """Export to CSV with utf-8-sig for proper Vietnamese characters in Excel."""
        if detailed:
            return self.export_detailed_csv(table_type=table_type, toxic_only=toxic_only, filename=filename, limit=limit)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "toxic_only" if toxic_only else "all"
        out_name = filename or f"threads_{table_type}_{suffix}_{timestamp}.csv"
        out_path = self.export_dir / out_name

        if table_type == "comments":
            df = db_manager.get_comments_df(toxic_only=toxic_only, limit=limit)
        else:
            df = db_manager.get_posts_df(toxic_only=toxic_only, limit=limit)

        df.to_csv(out_path, index=False, encoding="utf-8-sig")
        logger.info(f"Exported CSV file successfully: {out_path}")
        return out_path

    def export_to_excel(
        self,
        table_type: str = "comments",  # "comments" or "posts" or "both"
        toxic_only: bool = True,
        filename: Optional[str] = None,
        limit: Optional[int] = None,
        columns: Optional[List[str]] = None
    ) -> Path:
        """
        Export data to formatted Excel workbook with color highlights for toxic content.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "toxic_only" if toxic_only else "all"
        out_name = filename or f"threads_{table_type}_{suffix}_{timestamp}.xlsx"
        out_path = self.export_dir / out_name

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            if table_type in ["comments", "both"]:
                df_c = db_manager.get_detailed_export_df(table_type="comments", toxic_only=toxic_only, limit=limit)
                if columns:
                    valid_cols = [c for c in columns if c in df_c.columns]
                    if valid_cols:
                        df_c = df_c[valid_cols]
                df_c.to_excel(writer, sheet_name="Bình luận", index=False)

            if table_type in ["posts", "both"]:
                df_p = db_manager.get_detailed_export_df(table_type="posts", toxic_only=toxic_only, limit=limit)
                df_p.to_excel(writer, sheet_name="Bài đăng", index=False)

        self._format_excel_file(out_path)
        logger.info(f"Exported Excel file successfully: {out_path}")
        return out_path

    def export_curated_excel(
        self,
        filter_status: str = "all",
        filter_comment_type: str = "all",
        filename: Optional[str] = None,
        limit: Optional[int] = None,
        columns: Optional[List[str]] = None,
        viet_eng_only: bool = True
    ) -> Path:
        """
        Export curated data with fixed widths and wrap text.
        If columns is specified, only include selected visible columns.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_name = filename or f"threads_danh_gia_4_cot_{filter_status}_{timestamp}.xlsx"
        out_path = self.export_dir / out_name

        df = db_manager.get_curated_export_df(
            filter_status=filter_status,
            filter_comment_type=filter_comment_type,
            limit=limit,
            viet_eng_only=viet_eng_only
        )

        if columns:
            valid_cols = [c for c in columns if c in df.columns]
            if valid_cols:
                df = df[valid_cols]

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Dữ liệu đánh giá", index=False)

        self._format_curated_excel_file(out_path)
        logger.info(f"Exported curated Excel file successfully: {out_path} ({len(df)} rows, {len(df.columns)} cols)")
        return out_path

    def export_curated_csv(
        self,
        filter_status: str = "all",
        filter_comment_type: str = "all",
        filename: Optional[str] = None,
        limit: Optional[int] = None,
        columns: Optional[List[str]] = None,
        viet_eng_only: bool = True
    ) -> Path:
        """
        Export curated data to UTF-8-BOM CSV for Excel compatibility.
        If columns is specified, only include selected visible columns.
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_name = filename or f"threads_danh_gia_4_cot_{filter_status}_{timestamp}.csv"
        out_path = self.export_dir / out_name

        df = db_manager.get_curated_export_df(
            filter_status=filter_status,
            filter_comment_type=filter_comment_type,
            limit=limit,
            viet_eng_only=viet_eng_only
        )

        if columns:
            valid_cols = [c for c in columns if c in df.columns]
            if valid_cols:
                df = df[valid_cols]

        df.to_csv(out_path, index=False, encoding="utf-8-sig")
        logger.info(f"Exported curated CSV file successfully: {out_path} ({len(df)} rows, {len(df.columns)} cols)")
        return out_path

    def _format_curated_excel_file(self, file_path: Path):
        """
        Applies professional styling, fixed column widths, and text wrapping
        dynamically based on the header names of the columns present in the file.
        """
        import openpyxl
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Segoe UI", color="FFFFFF", bold=True, size=11)
        
        user_col_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        user_col_font = Font(name="Segoe UI", color="854D0E", bold=True, size=10)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Mapping of column configurations by header title
        COLUMN_CONFIG = {
            "Mã ID": {"width": 16, "align": "center", "font_size": 9.5, "color": "64748B", "is_user": False},
            "Nội dung": {"width": 48, "align": "left", "font_size": 10.5, "color": "000000", "is_user": False},
            "Điểm đánh giá": {"width": 22, "align": "center", "font_size": 10.5, "color": "000000", "is_user": False},
            "Bài viết": {"width": 48, "align": "left", "font_size": 10.5, "color": "000000", "is_user": False},
            "Chủ đề bài viết": {"width": 24, "align": "center", "font_size": 10.5, "color": "000000", "is_user": False},
            "Tự đánh giá của bạn (Xấu luôn / Chưa rõ / Trong sạch)": {"width": 30, "align": "center", "font_size": 10, "color": "854D0E", "is_user": True},
            "Từ lóng mới bổ sung (nếu có)": {"width": 28, "align": "center", "font_size": 10, "color": "854D0E", "is_user": True}
        }

        # Determine column config for each column by header name in row 1
        col_settings = {}
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header_val = str(ws.cell(row=1, column=col_idx).value or "").strip()
            # Match exact or partial
            matched_cfg = None
            for key, cfg in COLUMN_CONFIG.items():
                if key.lower() in header_val.lower() or header_val.lower() in key.lower():
                    matched_cfg = cfg
                    break
            if matched_cfg:
                col_settings[col_idx] = matched_cfg
                ws.column_dimensions[col_letter].width = matched_cfg["width"]
            else:
                col_settings[col_idx] = {"width": 25, "align": "left", "font_size": 10.5, "color": "000000", "is_user": False}
                ws.column_dimensions[col_letter].width = 25

        # ─── Xác định các cột cần tính chiều cao tự động ───
        # Cột "Nội dung" và "Bài viết" có thể rất dài → cần tự dãn chiều cao hàng
        LONG_TEXT_COL_KEYWORDS = ["nội dung", "bài viết"]
        long_text_col_indices = []
        for col_idx in range(1, ws.max_column + 1):
            header_val = str(ws.cell(row=1, column=col_idx).value or "").strip().lower()
            if any(kw in header_val for kw in LONG_TEXT_COL_KEYWORDS):
                long_text_col_indices.append(col_idx)

        # ─── Áp dụng style cho toàn bộ ô dữ liệu ───
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
                cfg = col_settings.get(cell.column, {"align": "left", "font_size": 10.5, "color": "000000", "is_user": False})
                cell.alignment = Alignment(horizontal=cfg["align"], vertical="top", wrap_text=True)
                if cfg.get("is_user"):
                    cell.fill = user_col_fill
                    cell.font = user_col_font
                else:
                    cell.font = Font(name="Segoe UI", size=cfg["font_size"], color=cfg.get("color", "000000"))

        # ─── Tự động tính chiều cao hàng theo nội dung dài nhất ───
        # Công thức: số dòng ≈ ceil(len(text) / col_width_in_chars)
        # Chiều cao 1 dòng trong Excel ≈ 14.4 pt (font size 10.5 pt)
        LINE_HEIGHT_PT = 14.4        # điểm chiều cao cho 1 dòng text
        MIN_ROW_HEIGHT = 18.0        # chiều cao tối thiểu (pt)
        MAX_ROW_HEIGHT = 400.0       # chiều cao tối đa để tránh hàng quá khổng lồ
        CHAR_WIDTH_CORRECTION = 1.8  # hệ số hiệu chỉnh vì ký tự Việt rộng hơn ASCII

        for row_idx in range(2, ws.max_row + 1):
            max_lines = 1
            for col_idx in long_text_col_indices:
                cell_val = ws.cell(row=row_idx, column=col_idx).value
                if not cell_val:
                    continue
                text = str(cell_val)
                # Số ký tự thực / chiều rộng cột (đơn vị ký tự Excel)
                col_width = col_settings.get(col_idx, {}).get("width", 48)
                # Tính số dòng cần dựa trên ký tự xuống dòng thực và chiều rộng cột
                lines_from_newlines = text.count("\n") + 1
                # Ước tính số dòng khi wrap: ký tự Việt/Unicode rộng hơn ASCII
                chars_per_line = max(1, int(col_width / CHAR_WIDTH_CORRECTION))
                # Đếm từng đoạn (paragraph) bị newline ngắt
                paragraphs = text.split("\n")
                wrap_lines = sum(
                    max(1, -(-len(p) // chars_per_line))  # ceiling division
                    for p in paragraphs
                )
                needed_lines = max(lines_from_newlines, wrap_lines)
                max_lines = max(max_lines, needed_lines)

            row_height = min(MAX_ROW_HEIGHT, max(MIN_ROW_HEIGHT, max_lines * LINE_HEIGHT_PT + 4))
            ws.row_dimensions[row_idx].height = row_height

        wb.save(file_path)


    def _format_excel_file(self, file_path: Path):
        """Format header, column widths, and highlight toxic rows in Excel."""
        import openpyxl
        wb = openpyxl.load_workbook(file_path)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        toxic_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # light red

        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for sheet in wb.worksheets:
            if sheet.max_row < 1:
                continue

            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            headers = [cell.value for cell in sheet[1]]
            toxic_col_idx = None
            for idx, h in enumerate(headers, 1):
                if h and ("Is Toxic" in str(h) or "Xúc phạm" in str(h) or h == "is_toxic"):
                    toxic_col_idx = idx
                    break

            for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
                is_toxic = False
                if toxic_col_idx:
                    val = row[toxic_col_idx - 1].value
                    if val in [1, True, "True", "Có", "1"]:
                        is_toxic = True

                for cell in row:
                    cell.border = thin_border
                    if is_toxic:
                        cell.fill = toxic_fill

            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

        wb.save(file_path)

    def export_to_json(
        self,
        table_type: str = "comments",
        toxic_only: bool = True,
        filename: Optional[str] = None,
        limit: Optional[int] = None
    ) -> Path:
        """Export to JSON for AI / NLP dataset training."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        suffix = "toxic_only" if toxic_only else "all"
        out_name = filename or f"threads_{table_type}_{suffix}_{timestamp}.json"
        out_path = self.export_dir / out_name

        if table_type == "comments":
            df = db_manager.get_comments_df(toxic_only=toxic_only, limit=limit)
        else:
            df = db_manager.get_posts_df(toxic_only=toxic_only, limit=limit)

        records = df.to_dict(orient="records")
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)

        logger.info(f"Exported JSON file successfully: {out_path}")
        return out_path

# Global exporter instance
exporter = DataExporter()
